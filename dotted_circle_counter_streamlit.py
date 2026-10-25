import io, math, hashlib
from datetime import datetime
from pathlib import Path
import cv2
import numpy as np
import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

SUPPORTED_EXTENSIONS=(".png",".jpg",".jpeg",".bmp",".tif",".tiff",".webp")
FILE_KEYWORD="no sew"
WORK_SIZE=2500; MIN_DIAMETER=8; MAX_DIAMETER_RATIO=.06; MIN_CIRCULARITY=.75; MIN_FILL_RATIO=.70
ASPECT_MIN,ASPECT_MAX=.8,1.25; DOT_CENTER_TOL=.35; DOT_MAX_AREA_RATIO=.35
HOLLOW_MIN_HOLE_RATIO=.10; HOLLOW_CENTER_TOL=.20; ISOLATION_MARGIN=1.0
MIN_SHAPE_AREA_RATIO=.005; VERTEX_ZONE=.22; EDGE_ZONE=.12

st.set_page_config(page_title="Dotted Circle Counter V11",page_icon="⭕",layout="wide",initial_sidebar_state="collapsed")
st.markdown("""<style>
.block-container{padding:12px 18px 10px;max-width:1900px}.stApp{font-family:'Segoe UI',sans-serif}
h1{font-size:26px!important;margin:0 0 4px!important}[data-testid=stSidebar],[data-testid=collapsedControl]{display:none}
[data-testid=stFileUploaderDropzone]{min-height:74px;border-radius:3px}div.stButton>button,div.stDownloadButton>button{border-radius:3px;min-height:38px}
[data-testid=stDataFrame]{border:1px solid #aaa}.preview{height:565px;background:#2b2b2b;border:1px solid #777;display:flex;align-items:center;justify-content:center;color:white}
.status{font-size:14px;padding-top:5px}.legend{padding-top:5px}.small-note{color:#777;font-size:12px}
</style>""",unsafe_allow_html=True)

for k,v in {"records":[],"selected":None,"signature":"","copy_text":""}.items():
    if k not in st.session_state: st.session_state[k]=v

def read_image(data):
    img=cv2.imdecode(np.frombuffer(data,np.uint8),cv2.IMREAD_COLOR)
    if img is None: raise ValueError("Không đọc được ảnh")
    return img

def children(h,i):
    out=[]; g=h[i][2]
    while g>=0: out.append(g); g=h[g][0]
    return out

def centroid(c):
    m=cv2.moments(c)
    return None if m["m00"]==0 else (m["m10"]/m["m00"],m["m01"]/m["m00"])

def find_shapes(contours,shape):
    H,W=shape; out=[]
    for c in contours:
        a=cv2.contourArea(c)
        if a<MIN_SHAPE_AREA_RATIO*H*W: continue
        p=cv2.approxPolyDP(c,.01*cv2.arcLength(c,True),True)
        if 3<=len(p)<=12: out.append((p.reshape(-1,2).astype(np.float32),a))
    return out

def isolated(bw,x,y,d):
    H,W=bw.shape; m=int(ISOLATION_MARGIN*d); half=int(d/2)+3
    x1,x2=max(0,int(x)-half-m),min(W,int(x)+half+m); y1,y2=max(0,int(y)-half-m),min(H,int(y)+half+m)
    win=bw[y1:y2,x1:x2].copy()
    if not win.size:return False
    cv2.circle(win,(int(x)-x1,int(y)-y1),half+1,0,-1)
    return np.count_nonzero(win)<.01*win.size

def near_vertex(x,y,r,shapes):
    for poly,area in shapes:
        size=math.sqrt(area)
        if 2*r>.25*size:continue
        de=cv2.pointPolygonTest(poly.reshape(-1,1,2),(float(x),float(y)),True)
        if de<=0 or de>EDGE_ZONE*size:continue
        if min(math.hypot(x-px,y-py) for px,py in poly)<=VERTEX_ZONE*size:return True
    return False

def depth(h,i):
    d=0;p=h[i][3]
    while p>=0:d+=1;p=h[p][3]
    return d

def dedup(items):
    out=[]
    for f in sorted(items,key=lambda t:-t[2]):
        if all(math.hypot(f[0]-o[0],f[1]-o[1])>max(f[2],o[2]) for o in out):out.append(f)
    return out

def detect_circles(img):
    gray=cv2.cvtColor(img,cv2.COLOR_BGR2GRAY); h,w=gray.shape; scale=1.
    if max(h,w)<WORK_SIZE:
        scale=WORK_SIZE/float(max(h,w)); gray=cv2.resize(gray,None,fx=scale,fy=scale,interpolation=cv2.INTER_CUBIC)
    _,bw=cv2.threshold(gray,0,255,cv2.THRESH_BINARY_INV+cv2.THRESH_OTSU)
    contours,hier=cv2.findContours(bw,cv2.RETR_TREE,cv2.CHAIN_APPROX_SIMPLE)
    if hier is None:return []
    hier=hier[0]; maxd=MAX_DIAMETER_RATIO*max(bw.shape); shapes=find_shapes(contours,bw.shape)
    mask=np.zeros_like(bw); lw=max(3,int(.004*max(bw.shape)))
    for c in contours:
        if cv2.contourArea(c)>=MIN_SHAPE_AREA_RATIO*bw.size:cv2.drawContours(mask,[c],-1,255,lw)
    noframe=cv2.bitwise_and(bw,cv2.bitwise_not(mask)); dotted=[]; hollow=[]
    for i,c in enumerate(contours):
        if depth(hier,i)%2==1:continue
        kids=children(hier,i)
        if len(kids)!=1:continue
        bx,by,cw,ch=cv2.boundingRect(c)
        if cw<MIN_DIAMETER or ch<MIN_DIAMETER or cw>maxd or ch>maxd or not ASPECT_MIN<=cw/float(ch)<=ASPECT_MAX:continue
        x,y=bx+cw/2.,by+ch/2.; r=(cw+ch)/4.; area=cv2.contourArea(c); per=cv2.arcLength(c,True)
        if per==0 or 4*math.pi*area/(per*per)<MIN_CIRCULARITY or area/(math.pi*r*r)<MIN_FILL_RATIO:continue
        hole=kids[0]; ha=cv2.contourArea(contours[hole]); hc=centroid(contours[hole])
        if hc is None or math.hypot(hc[0]-x,hc[1]-y)>HOLLOW_CENTER_TOL*r:continue
        inner=children(hier,hole)
        if len(inner)==1 and ha>.25*area:
            dc=centroid(contours[inner[0]])
            if dc and math.hypot(dc[0]-x,dc[1]-y)<DOT_CENTER_TOL*r and cv2.contourArea(contours[inner[0]])<DOT_MAX_AREA_RATIO*ha:dotted.append((x,y,r))
        elif not inner and ha>=HOLLOW_MIN_HOLE_RATIO*area:
            if isolated(noframe,x,y,2*r) and near_vertex(x,y,r,shapes):hollow.append((x,y,r))
    dotted=dedup(dotted)
    hollow=[h for h in dedup(hollow) if all(math.hypot(h[0]-d[0],h[1]-d[1])>max(h[2],d[2])*1.5 for d in dotted)]
    result=[(x,y,r,"dot") for x,y,r in dotted]+[(x,y,r,"hollow") for x,y,r in hollow]
    return sorted([(int(round(x/scale)),int(round(y/scale)),max(2,int(round(r/scale))),k) for x,y,r,k in result],key=lambda t:(t[1],t[0]))

def draw_result(img,circles):
    out=img.copy()
    for i,(x,y,r,kind) in enumerate(circles,1):
        color=(0,200,0) if kind=="dot" else (0,140,255)
        cv2.circle(out,(x,y),r+6,color,2,cv2.LINE_AA);cv2.putText(out,str(i),(x+r+8,max(18,y-r-4)),cv2.FONT_HERSHEY_SIMPLEX,.6,(0,0,255),2,cv2.LINE_AA)
    out=cv2.copyMakeBorder(out,50,0,0,0,cv2.BORDER_CONSTANT,value=(255,255,255));cv2.putText(out,f"TOTAL: {len(circles)}",(10,36),cv2.FONT_HERSHEY_SIMPLEX,.9,(0,0,255),2,cv2.LINE_AA)
    ok,b=cv2.imencode('.png',out)
    if not ok:raise ValueError("Không tạo được ảnh kết quả")
    return b.tobytes()

def get_files(uploaded):
    out=[];seen=set()
    for f in uploaded or []:
        name=str(f.name).replace('\\','/'); base=Path(name).name
        if FILE_KEYWORD not in base.lower() or Path(base).suffix.lower() not in SUPPORTED_EXTENSIONS:continue
        data=f.getvalue(); sig=hashlib.sha256(data).hexdigest()
        if sig in seen:continue
        seen.add(sig);out.append({"name":base,"path":name,"data":data,"hash":sig})
    return sorted(out,key=lambda x:x['name'].lower())

def excel_bytes(records):
    wb=Workbook();ws=wb.active;ws.title="Ket qua";ws.append(["STT","Tên bản vẽ","Số lượng","Trạng thái","Đường dẫn bản vẽ"])
    thin=Side(style="thin",color="BFBFBF");border=Border(left=thin,right=thin,top=thin,bottom=thin)
    for c in ws[1]:c.font=Font(bold=True,color="FFFFFF");c.fill=PatternFill("solid",fgColor="1F4E78");c.alignment=Alignment(horizontal="center");c.border=border
    for i,r in enumerate(records,1):ws.append([i,r['name'],r['count'] if r['ok'] else None,r['status'],r['path']])
    last=ws.max_row;ws.append(["","TỔNG",f"=SUM(C2:C{last})","",""])
    for c in ws[ws.max_row]:c.font=Font(bold=True);c.fill=PatternFill("solid",fgColor="DDEBF7")
    for row in ws.iter_rows(min_row=2,max_row=ws.max_row):
        for c in row:c.border=border
    for col,width in zip("ABCDE",(7,45,11,18,70)):ws.column_dimensions[col].width=width
    ws.freeze_panes="A2";ws.auto_filter.ref=f"A1:E{last}";buf=io.BytesIO();wb.save(buf);return buf.getvalue()

def process(files):
    records=[];bar=st.progress(0);message=st.empty()
    for i,f in enumerate(files,1):
        message.info(f"Đang xử lý {i}/{len(files)}: {f['name']}")
        r={**f,"count":0,"circles":[],"status":"","ok":False,"preview":b""}
        try:
            img=read_image(f['data']);r['circles']=detect_circles(img);r['count']=len(r['circles']);r['preview']=draw_result(img,r['circles']);r['status']="Hoàn tất";r['ok']=True
        except Exception as e:r['status']=f"Lỗi: {e}"
        records.append(r);bar.progress(i/len(files))
    message.empty();return records

st.title("ĐẾM VÒNG TRÒN TRÊN BẢN VẼ NO SEW")
try:
    uploaded=st.file_uploader("Folder bản vẽ:",type=[x[1:] for x in SUPPORTED_EXTENSIONS],accept_multiple_files="directory")
except TypeError:
    uploaded=st.file_uploader("Folder bản vẽ:",type=[x[1:] for x in SUPPORTED_EXTENSIONS],accept_multiple_files=True)
files=get_files(uploaded); signature='|'.join(x['hash'] for x in files)

c1,c2,c3,c4=st.columns([1,1,1,6])
run=c1.button("Phân tích",type="primary",disabled=not files,use_container_width=True)
copy=c2.button("Copy bảng",disabled=not st.session_state.records,use_container_width=True)
if st.session_state.records:
    c3.download_button("Xuất Excel",excel_bytes(st.session_state.records),f"NoSew_Circle_Count_{datetime.now():%Y%m%d_%H%M%S}.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True)
else:c3.button("Xuất Excel",disabled=True,use_container_width=True)
if run:
    st.session_state.records=process(files);st.session_state.signature=signature;st.session_state.selected=st.session_state.records[0]['path'] if st.session_state.records else None;st.rerun()
if copy:
    lines=["STT\tTên bản vẽ\tSố lượng\tTrạng thái"]+[f"{i}\t{r['name']}\t{r['count'] if r['ok'] else ''}\t{r['status']}" for i,r in enumerate(st.session_state.records,1)]
    lines.append(f"\tTỔNG\t{sum(r['count'] for r in st.session_state.records if r['ok'])}\t");st.code('\n'.join(lines),language=None)

records=st.session_state.records; total=sum(r['count'] for r in records if r['ok'])
left,right=st.columns([2,3],gap="medium")
with left:
    a,b=st.columns(2);a.markdown(f"**Bản vẽ: {len(files) if files else len(records)}**");b.markdown(f"<div style='text-align:right'><b>Tổng ký hiệu: {total}</b></div>",unsafe_allow_html=True)
    rows=[]
    source=records if records else files
    for i,r in enumerate(source,1):rows.append({"STT":i,"Tên bản vẽ":r['name'],"Số lượng":r.get('count'),"Trạng thái":r.get('status','Chưa chạy')})
    st.dataframe(pd.DataFrame(rows,columns=["STT","Tên bản vẽ","Số lượng","Trạng thái"]),hide_index=True,use_container_width=True,height=565,column_config={"STT":st.column_config.NumberColumn(width="small"),"Tên bản vẽ":st.column_config.TextColumn(width="large"),"Số lượng":st.column_config.NumberColumn(width="small"),"Trạng thái":st.column_config.TextColumn(width="medium")})
with right:
    h1,h2=st.columns([1,3]);h1.markdown("**Xem bản vẽ**");selected=None
    if records:
        opts=[r['path'] for r in records];idx=opts.index(st.session_state.selected) if st.session_state.selected in opts else 0
        choice=h2.selectbox("Bản vẽ",opts,index=idx,label_visibility="collapsed");st.session_state.selected=choice;selected=next(r for r in records if r['path']==choice)
        h2.caption(f"{selected['name']} | Số lượng: {selected['count'] if selected['ok'] else ''}")
    box=st.container(height=565,border=True)
    with box:
        if selected and selected['ok']:st.image(selected['preview'],use_container_width=True)
        elif selected:st.error(selected['status'])
        else:st.markdown("<div style='text-align:center;padding-top:230px;color:#888'>Bấm Phân tích rồi chọn tên bản vẽ để xem hình và kết quả</div>",unsafe_allow_html=True)
    st.markdown("<div class='legend'><span style='color:#00C800'>●</span> Vòng có chấm &nbsp;&nbsp;&nbsp; <span style='color:#FF8C00'>●</span> Vòng rỗng</div>",unsafe_allow_html=True)

if records:st.progress(1.0);status=f"Hoàn tất {len(records)} bản vẽ. Tổng ký hiệu: {total}. Bấm vào tên bản vẽ để xem."
elif files:st.progress(0.0);status=f"Tìm thấy {len(files)} bản vẽ có tên chứa NO SEW."
else:st.progress(0.0);status="Chọn folder chứa bản vẽ để bắt đầu."
st.markdown(f"<div class='status'>{status}</div>",unsafe_allow_html=True)
